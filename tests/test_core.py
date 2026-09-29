"""Offline tests: no network, no API calls. A fake LLM stands in for the engine."""
import json

import pytest
from openpyxl import load_workbook
from pptx import Presentation

from ca_tool import config, store
from ca_tool.export import excel, pptx
from ca_tool.knowledge.scoring import SCORING_GUIDE, guide_text, label
from ca_tool.knowledge.sources import classify
from ca_tool.knowledge.template import ALL_METRICS, PRICING
from ca_tool.llm.base import LLM
from ca_tool.models import (Candidate, Cell, Citation, CompetitorRecord, FeatureDef, FeatureScore, Intake,
                            Project, Quote, Source)
from ca_tool.pipeline import collect, score, synthesize
from ca_tool.research import web
from ca_tool.research.corpus import Corpus, quote_in_text

PAGE = ("Acme Spaces was founded in 2015 and is headquartered in Austin, Texas. "
        "Our platform lets venues build interactive 3D floor plans and share virtual tours with planners. "
        "“Setup was quick and the support team answered within an hour,” says one G2 reviewer. "
        "Another reviewer wrote that the mobile app crashes when loading large floor plans.")


class FakeLLM(LLM):
    """Answers each schema with deterministic data citing S1."""
    name = "fake"
    last_model = "fake"

    def generate_text(self, prompt, system=""):
        return "ok"

    def generate_json(self, prompt, schema, system=""):
        n = schema.__name__
        cite = [{"source_id": "S1", "quote": "founded in 2015 and is headquartered in Austin, Texas"}]
        if n == "_CollectOut":
            keys = [m.key for m in ALL_METRICS + PRICING if f"- {m.key}:" in prompt]
            vals = [{"key": k, "value": f"value for {k}", "confidence": "high", "note": "", "citations": cite} for k in keys]
            vals[-1] = {"key": keys[-1], "value": "Undisclosed", "confidence": "unavailable", "note": "use proxy", "citations": []}
            data = {"values": vals}
        elif n == "_ScoreOut":
            data = {"scores": [
                {"id": "F1", "feature": "3D floor plans", "standard": True, "premium": True, "enterprise": None, "score": 4,
                 "rationale": "Advanced", "confidence": "high",
                 "citations": [{"source_id": "S1", "quote": "build interactive 3D floor plans"}]},
                {"id": "F2", "feature": "VR headset support", "standard": None, "premium": None, "enterprise": None, "score": None,
                 "rationale": "no evidence", "confidence": "unavailable", "citations": []},
                # hallucinated citation: quote not on the page
                {"id": "X", "feature": "[Planning] Seating charts", "standard": True, "premium": True, "enterprise": True, "score": 5,
                 "rationale": "Leading", "confidence": "high",
                 "citations": [{"source_id": "S1", "quote": "industry leading seating chart engine used by millions"}]},
            ]}
        elif n == "_CompOut":
            pt = [{"text": "Strong 3D", "citations": [{"source_id": "S1", "quote": "build interactive 3D floor plans"}]}]
            data = {"strengths": pt, "weaknesses": pt, "opportunities": pt, "threats": pt,
                    "voice_of_customer": [
                        {"text": "Setup was quick and the support team answered within an hour", "sentiment": "positive", "source_id": "S1"},
                        {"text": "Invented quote that nobody ever wrote anywhere at all", "sentiment": "negative", "source_id": "S1"}],
                    "narrative": "Acme is winning in venues because of 3D.", "how_client_wins": ["Lead with seating"]}
        elif n == "_InsightsOut":
            data = {k: "• point" for k in ("executive_summary", "white_space", "trends", "implications")}
        else:
            raise AssertionError(n)
        return schema.model_validate_json(json.dumps(data))


@pytest.fixture
def tmp_projects(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "PROJECTS_DIR", tmp_path / "projects")
    return tmp_path


@pytest.fixture
def corpus():
    c = Corpus()
    c.add("https://acme.example/about", "About Acme", PAGE, kind="website")
    return c


def test_quote_verification_exact_and_fuzzy():
    assert quote_in_text("founded in 2015 and is headquartered in Austin", PAGE)
    assert quote_in_text("Setup was quick and the support team answered within an hour", PAGE)  # curly quotes on page
    assert quote_in_text("FOUNDED in 2015   and is headquartered in Austin, Texas.", PAGE)  # case/space
    assert not quote_in_text("founded in 2012 and is headquartered in Denver, Colorado", PAGE)
    assert not quote_in_text("short", PAGE)


def test_quote_verification_bullets_and_ellipses():
    page = ("Its ecosystem has expanded into\n- commercial office\n- hospitality\n- retail\n"
            "Enterprise customers mention account managers by name, including Saba and Dan, "
            "and describe them as responsive and proactive.")
    assert quote_in_text("expanded into: commercial office, hospitality, retail", page)
    assert quote_in_text("Enterprise customers mention account managers by name... and describe them as responsive", page)
    assert quote_in_text("Enterprise customers mention account managers by name … responsive and proactive", page)
    # parts out of order or a fabricated part must fail
    assert not quote_in_text("describe them as responsive... Enterprise customers mention account managers", page)
    assert not quote_in_text("Enterprise customers mention account managers by name... and pay $9.99 per month", page)
    # page titles are not evidence
    assert not quote_in_text("Cvent | Event Platform for In-person, Virtual and Hybrid Events", "Welcome to our site")


def test_corpus_resolve_drops_unknown_sources_and_flags_unverified(corpus):
    out = corpus.resolve([Citation(source_id="S1", quote="build interactive 3D floor plans"),
                          Citation(source_id="[S1]", quote="completely made up claim about pricing tiers"),
                          Citation(source_id="S9", quote="anything")])
    assert [c.verified for c in out] == [True, False]
    assert out[0].url == "https://acme.example/about"


def test_corpus_dedupes_and_respects_budget(corpus):
    assert corpus.add("https://acme.example/about", "dup", "x" * 500) is None
    corpus.add("https://b.example", "B", "y" * 5000)
    assert len(corpus.render(max_chars=1000)) <= 1010


def test_source_priorities():
    assert classify("https://www.g2.com/products/x/reviews") == ("review", 1)
    assert classify("https://news.crunchbase.com/a") == ("funding", 1)
    assert classify("https://pinterest.com/pin/1")[1] == 9
    hits = [web.SearchHit("a", "https://blog.example/x", ""), web.SearchHit("b", "https://www.g2.com/x", ""),
            web.SearchHit("c", "https://pinterest.com/y", ""), web.SearchHit("d", "https://www.g2.com/x", "")]
    assert [h.url for h in web.rank_hits(hits)] == ["https://www.g2.com/x", "https://blog.example/x"]


def test_scoring_guide_is_complete():
    assert sorted(SCORING_GUIDE) == [0, 1, 2, 3, 4, 5]
    assert all(set(d) == {"Presence", "Quality", "Maturity"} for _, d in SCORING_GUIDE.values())
    assert label(None) == "?" and label(4.0) == "Advanced"
    assert "5 - Leading" in guide_text()


def test_key_pages_and_guessed_paths():
    home = web.Page(url="https://acme.example/", title="", text="", links=[
        ("About us", "https://acme.example/company/about"), ("Pricing", "https://acme.example/pricing"),
        ("Partner site", "https://other.example/partners")])
    kp = web.key_pages(home)
    assert kp["about"].endswith("/company/about") and kp["pricing"].endswith("/pricing") and "partners" not in kp
    g = web.guessed_pages("https://acme.example/", ["careers"])
    assert ("careers", "https://acme.example/careers") in g


def test_extract_fills_every_metric_with_verified_citations(corpus):
    intake = Intake(client_name="Client", product="3D venue software")
    cells = collect.extract("Acme", intake, corpus, ALL_METRICS + PRICING, FakeLLM())
    assert set(cells) == {m.key for m in ALL_METRICS + PRICING}
    assert cells["founding_year"].verified and cells["founding_year"].confidence == "high"
    last = (ALL_METRICS + PRICING)[-1].key
    assert cells[last].confidence == "unavailable" and cells[last].note == "use proxy"


def _record(corpus) -> CompetitorRecord:
    return CompetitorRecord(name="Acme", website="https://acme.example", sources=corpus.sources)


def test_scoring_unknown_is_question_mark_and_fake_quotes_downgraded(corpus):
    feats = [FeatureDef(category="Design", name="3D floor plans"), FeatureDef(category="Design", name="VR headset support"),
             FeatureDef(category="Planning", name="Seating charts"), FeatureDef(category="Planning", name="Budgeting")]
    out = {f.feature: f for f in score.score(_record(corpus), feats, Intake(client_name="C", product="p"), FakeLLM())}
    assert out["3D floor plans"].score == 4 and out["3D floor plans"].confidence == "high"
    assert out["VR headset support"].score is None
    assert out["Seating charts"].confidence == "low"          # quote not found on page
    assert out["Budgeting"].confidence == "unavailable"       # model skipped it


def test_synthesis_keeps_only_verified_customer_quotes(corpus):
    rec = _record(corpus)
    synthesize.synthesize_competitor(rec, Intake(client_name="C", product="p"), FakeLLM())
    assert set(rec.swot) == {"strengths", "weaknesses", "opportunities", "threats"}
    assert [q.sentiment for q in rec.voc] == ["positive"]      # invented negative quote dropped
    assert rec.battlecard.value.startswith("•")


def _full_project(corpus) -> Project:
    p = Project(id="t", intake=Intake(client_name="Client", product="3D venue software"))
    p.features = [FeatureDef(category="Design", name="3D floor plans"), FeatureDef(category="Planning", name="Seating charts")]
    llm = FakeLLM()
    for name in ("Acme", "Beta"):
        rec = CompetitorRecord(name=name, website=f"https://{name.lower()}.example", sources=corpus.sources)
        cells = collect.extract(name, p.intake, corpus, ALL_METRICS + PRICING, llm)
        rec.pricing = {k: cells.pop(k) for k in collect.PRICING_BY_KEY}
        rec.cells = cells
        rec.features = score.score(rec, p.features, p.intake, llm)
        synthesize.synthesize_competitor(rec, p.intake, llm)
        p.competitors.append(rec)
    p.competitors[1].cells["description"].status = "rejected"
    p.insights = synthesize.synthesize_insights(p, llm)
    return p


def test_excel_export_matches_template_layout(tmp_path, corpus):
    p = _full_project(corpus)
    path = excel.export(p, tmp_path / "out.xlsx")
    wb = load_workbook(path)
    assert {"Company Analysis", "Feature Comparison", "Scoring Guide", "Insights", "Sources"} <= set(wb.sheetnames)
    ca = wb["Company Analysis"]
    assert ca["B2"].value == "Metric" and ca["C2"].value == "What it means?" and ca["D2"].value == "Acme"
    assert ca["D3"].hyperlink is not None
    row = next(r for r in range(4, 60) if ca.cell(r, 2).value == "Founding Year")
    assert ca.cell(row, 4).comment and "✓" in ca.cell(row, 4).comment.text
    desc = next(r for r in range(4, 60) if ca.cell(r, 2).value == "Description")
    assert ca.cell(desc, 5).fill.fgColor.rgb.endswith("F4CCCC")  # rejected -> red
    fc = wb["Feature Comparison"]
    assert fc["C6"].value == "Standard" and fc["F6"].value == "Score" and fc["G6"].value == "Context Notes"
    assert fc["F7"].value == 4 and fc["C5"].value == "Acme"


def test_pptx_exports(tmp_path, corpus):
    p = _full_project(corpus)
    deck = Presentation(pptx.export_deck(p, tmp_path / "deck.pptx"))
    titles = [next((sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text), "") for s in deck.slides]
    assert "Executive Summary" in titles and "Acme Overview" in titles and "Acme – SWOT Analysis" in titles
    assert any(t.startswith("Voice of the Customer – Acme") for t in titles)
    body = " ".join(sh.text_frame.text for s in deck.slides for sh in s.shapes if sh.has_text_frame)
    assert "Source: acme.example" in body
    cards = Presentation(pptx.export_battlecards(p, tmp_path / "cards.pptx"))
    assert len(cards.slides) == 2


def test_store_roundtrip(tmp_projects):
    p = store.create(Intake(client_name="Demo Co!", product="x", known_competitors=["A"]))
    assert p.id == "demo-co"
    p.candidates.append(Candidate(name="A", citations=[Citation(source_id="D1", quote="q", verified=True)]))
    store.save(p)
    q = store.load("demo-co")
    assert q.candidates[0].citations[0].verified and store.list_ids() == ["demo-co"]
    assert store.create(Intake(client_name="Demo Co!", product="x")).id == "demo-co-2"


def test_scoring_fails_loudly_when_output_unmatchable(corpus):
    class Garbled(FakeLLM):
        def generate_json(self, prompt, schema, system=""):
            out = super().generate_json(prompt, schema, system)
            for sc in getattr(out, "scores", []):
                sc.id, sc.feature = "?", "something else entirely"
            return out
    feats = [FeatureDef(category="Design", name="3D floor plans"), FeatureDef(category="Design", name="VR headset support")]
    with pytest.raises(ValueError):
        score.score(_record(corpus), feats, Intake(client_name="C", product="p"), Garbled())


def test_deck_text_cleaning():
    out = pptx.clean("• Consolidation [S11], signal. • AI adoption [S14] and [S1, S9] models. • Third [D3]")
    assert out == "• Consolidation, signal.\n• AI adoption and models.\n• Third"


def test_snippet_citation_resolves_to_real_url():
    from ca_tool.research.corpus import snippet_url
    c = Corpus()
    c.add("search://X", "results", "- G2: I find the tool incredibly helpful for reducing admin tasks (https://www.g2.com/x)\n"
                                   "- Other page: unrelated text here (https://a.example/c)")
    cite = c.resolve([Citation(source_id="S1", quote="I find the tool incredibly helpful")])[0]
    assert cite.verified and cite.url == "https://www.g2.com/x"
    assert snippet_url("no urls here", "anything at all goes") == ""


def test_company_mention_matching():
    assert web.mentions("Try Event Diagram today", "EventDiagram")
    assert web.mentions("see eventdiagram.com for pricing", "EventDiagram")
    assert web.mentions("Cvent's diagramming tool", "Cvent Event Diagramming")
    assert not web.mentions("A page about floor plans", "EventDiagram")


def test_blocked_review_sites_are_snippet_only():
    from ca_tool.knowledge.sources import fetchable, site_name
    assert site_name("https://www.capterra.co.uk/x") == site_name("https://capterra.ae/y") == "capterra"
    assert not fetchable("https://www.g2.com/products/x/reviews")
    assert fetchable("https://en.wikipedia.org/wiki/Matterport")


def test_llm_waits_and_retries_when_all_models_busy(monkeypatch):
    from ca_tool.llm import gemini as g
    from google.genai import errors

    calls = []

    class Models:
        def generate_content(self, model, contents, config):
            calls.append(model)
            if len(calls) <= 2:
                raise errors.APIError(503, {"error": {"code": 503, "message": "busy", "status": "UNAVAILABLE"}})
            part = type("P", (), {"text": "ok"})()
            cand = type("C", (), {"content": type("X", (), {"parts": [part]})()})()
            return type("R", (), {"candidates": [cand]})()

    llm = g.GeminiLLM(api_key="x", models=["m1", "m2"], retries=1, round_wait=0)
    llm.client = type("Cl", (), {"models": Models()})()
    g.GeminiLLM._cooldown.clear()
    assert llm.generate_text("hi") == "ok" and calls == ["m1", "m2", "m1"]
    g.GeminiLLM._cooldown.clear()


def test_two_pass_discovery(monkeypatch):
    """Segments -> first pass -> 'alternatives to X' pass adds only new vendors; client never listed."""
    from ca_tool.pipeline import discover as d

    pages = {"https://list.example/a": "Top tools: Acme Spaces builds interactive 3D floor plans for venues.",
             "https://alt.example/b": "Alternatives to Acme include Niche Venue Co for hotels and venues."}
    monkeypatch.setattr(d, "search", lambda q, max_results=8: [
        web.SearchHit("t", "https://alt.example/b" if "alternatives competitors" in q else "https://list.example/a", "snippet")])
    monkeypatch.setattr(d, "fetch", lambda u: web.Page(url=u, title="p", text=pages[u] * 20, links=[]))

    class DiscLLM(FakeLLM):
        def generate_json(self, prompt, schema, system=""):
            n = schema.__name__
            if n == "_SegmentsOut":
                data = {"segments": [{"name": "3D diagramming", "why": "same buyers", "queries": ["best 3d tools"]}]}
            elif n == "_DiscoveryOut" and "Already listed" in prompt:
                sid = re.search(r"\[(D\d+)\][^\n]*\nURL: https://alt\.example/b", prompt).group(1)
                data = {"candidates": [
                    {"name": "Acme Spaces", "website": "", "type": "direct", "segment": "3D diagramming", "justification": "dup",
                     "size_group": "", "geography": "", "citations": []},
                    {"name": "Niche Venue Co", "website": "https://niche.example", "type": "direct", "segment": "3D diagramming",
                     "justification": "specialist", "size_group": "", "geography": "",
                     "citations": [{"source_id": sid, "quote": "Alternatives to Acme include Niche Venue Co"}]}]}
            elif n == "_DiscoveryOut":
                data = {"candidates": [
                    {"name": "Acme Spaces", "website": "https://acme.example", "type": "direct", "segment": "3D diagramming",
                     "justification": "3D floor plans", "size_group": "", "geography": "",
                     "citations": [{"source_id": "D2", "quote": "Acme Spaces builds interactive 3D floor plans"}]},
                    {"name": "Client", "website": "", "type": "direct", "segment": "", "justification": "self",
                     "size_group": "", "geography": "", "citations": []}]}
            else:
                return super().generate_json(prompt, schema, system)
            return schema.model_validate_json(json.dumps(data))

    import re
    p = Project(id="d", intake=Intake(client_name="Client", product="3D venue software", target_markets="hotels"))
    out = d.find_competitors(p, DiscLLM())
    assert [c.name for c in out] == ["Acme Spaces", "Niche Venue Co"]
    assert out[0].segment == "3D diagramming" and out[0].citations[0].verified
    assert out[1].citations and out[1].citations[0].verified


def test_market_map_slide(tmp_path, corpus):
    p = _full_project(corpus)
    p.candidates = [Candidate(name="Acme", segment="3D diagramming", status="accepted"),
                    Candidate(name="Beta", segment="Virtual tours", status="accepted"),
                    Candidate(name="Gamma", segment="Virtual tours", status="accepted"),
                    Candidate(name="Delta", segment="Venue booking", status="rejected")]
    deck = Presentation(pptx.export_deck(p, tmp_path / "deck.pptx"))
    slide = next(s for s in deck.slides if any(sh.has_text_frame and sh.text_frame.text == "Competitive landscape by market segment" for sh in s.shapes))
    texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
    assert {"3D diagramming", "Virtual tours", "Acme", "Gamma"} <= set(texts)
    assert "Delta" not in texts and "Venue booking" not in texts   # rejected at Gate 1


def test_client_aliases_are_excluded():
    from ca_tool.pipeline.discover import client_aliases
    c = Corpus(prefix="D")
    c.add("https://g2.example/x", "t", "Prismm (formerly Allseated) is a world leader in spatial design. Other text " * 3)
    c.add("https://b.example/y", "t", "Acme, formerly known as Beta Tools, sells floor plans. " * 3)
    assert client_aliases(c, "Prismm") == {"allseated"}
    assert client_aliases(c, "Acme") == {"beta tools"}
    assert client_aliases(c, "") == set()


def test_unsupported_zero_becomes_question_mark(corpus):
    class Zero(FakeLLM):
        def generate_json(self, prompt, schema, system=""):
            data = {"scores": [
                {"id": "F1", "feature": "Budgeting", "standard": None, "premium": None, "enterprise": None, "score": 0,
                 "rationale": "not mentioned", "confidence": "medium", "citations": []},
                {"id": "F2", "feature": "3D floor plans", "standard": False, "premium": False, "enterprise": False, "score": 0,
                 "rationale": "absent", "confidence": "high",
                 "citations": [{"source_id": "S1", "quote": "build interactive 3D floor plans"}]}]}
            return schema.model_validate_json(json.dumps(data))
    feats = [FeatureDef(category="Planning", name="Budgeting"), FeatureDef(category="Design", name="3D floor plans")]
    out = {f.feature: f for f in score.score(_record(corpus), feats, Intake(client_name="C", product="p"), Zero())}
    assert out["Budgeting"].score is None and "needs checking" in out["Budgeting"].rationale
    assert out["3D floor plans"].score == 0          # a zero backed by verified evidence stands


def test_pilot_fact_matching_is_strict():
    from ca_tool.pilot import fact_agrees as f, same_company
    assert f("headquarters", "Portland, Oregon US", "Carlton, Oregon, United States") is False   # same state, other city
    assert f("headquarters", "Tysons, Virginia, US", "Tysons, Virginia, USA") is True
    assert f("headquarters", "New York, US", "United States") is False                           # less specific
    assert f("employees", "11-50", "51-200") is False                                            # adjacent bands differ
    assert f("employees", "500-1000", "501-1000") is True
    assert f("employees", "3,601", "5,500+") is False
    assert f("founding_year", "2009", "2011") is False
    assert f("founding_year", "2011", "Undisclosed") is None
    assert same_company("Social Tables (Acquired by Cvent)", "Social Tables") and not same_company("Cvent", "Merri")


def test_figures_must_appear_in_the_evidence():
    from ca_tool.research.corpus import numbers_supported
    assert numbers_supported("2011", ["Matterport pioneered 3D digital twins"]) is False      # real quote, no year
    assert numbers_supported("Founded in 2011", ["Matterport, founded in 2011, ..."]) is True
    assert numbers_supported("Solo ($9/mo), Pro ($18/mo)", ["Solo Plan $9.00... Pro Plan $18.00"]) is True
    assert numbers_supported("3,601 employees", ["a team of 3601 people"]) is True
    assert numbers_supported("Worldwide", ["global presence"]) is None                       # nothing to check


def test_extract_downgrades_unsupported_figure(corpus):
    class Mem(FakeLLM):
        def generate_json(self, prompt, schema, system=""):
            data = {"values": [{"key": "founding_year", "value": "2012", "confidence": "high", "note": "",
                                "citations": [{"source_id": "S1", "quote": "build interactive 3D floor plans"}]},
                               {"key": "headquarters", "value": "Austin, Texas", "confidence": "high", "note": "",
                                "citations": [{"source_id": "S1", "quote": "headquartered in Austin, Texas"}]}]}
            return schema.model_validate_json(json.dumps(data))
    from ca_tool.knowledge.template import METRIC_BY_KEY
    cells = collect.extract("Acme", Intake(client_name="C", product="p"), corpus,
                            [METRIC_BY_KEY["founding_year"], METRIC_BY_KEY["headquarters"]], Mem())
    assert cells["founding_year"].verified and cells["founding_year"].confidence == "low"
    assert "does not appear in the quoted evidence" in cells["founding_year"].note
    assert cells["headquarters"].confidence == "high"


def test_quote_checker_rejects_adversarial_quotes():
    page = ("Acme offers a free plan for small businesses. Our support team is available by email "
            "and the company raised money from investors in a round led by well known firms last spring.")
    # invented ending after a real opening
    assert not quote_in_text("Acme offers a free plan for enterprises with unlimited SSO", page)
    # tiny fragments joined by ellipses, matching inside other words
    assert not quote_in_text("Best ... in ... class ... support", "The best tools in the classroom need support staff")
    # an invented middle section in a long quote
    assert not quote_in_text("the company raised 500 million dollars from investors in a round led by well known firms", page)
    # no word may differ, even in a long quote (one-word edits flipped meaning in review)
    assert not quote_in_text("the company raised funds from investors in a round led by well known firms last spring", page)
    assert quote_in_text("the company raised money from investors in a round led by well known firms last spring", page)
    # whole words only
    assert not quote_in_text("free plan for small business", page.replace("businesses", "businesspeople"))


def test_search_bundle_quotes_cannot_be_stitched():
    c = Corpus()
    c.add("search://X", "results", "- A: Acme builds interactive floor plans for venues (https://a.example/1)\n"
                                   "- B: Beta raised forty million dollars last year (https://b.example/2)")
    stitched = c.resolve([Citation(source_id="S1", quote="Acme builds interactive floor plans for venues B Beta raised forty million")])[0]
    single = c.resolve([Citation(source_id="S1", quote="Beta raised forty million dollars last year")])[0]
    assert not stitched.verified
    assert single.verified and single.url == "https://b.example/2"


def test_short_quotes_need_a_figure():
    page = "Company profile. Founded: 2011. Employees: 51-200. Standard onboarding is included."
    assert quote_in_text("Founded: 2011", page) and quote_in_text("Employees: 51-200", page)
    assert not quote_in_text("Founded: 2012", page)
    assert not quote_in_text("Standard onboarding", page)      # short and no figure: could match anywhere
    assert not quote_in_text("2011", page)


def test_excel_is_safe_against_formulas_and_control_chars(tmp_path, corpus):
    p = _full_project(corpus)
    rec = p.competitors[0]
    rec.cells["description"].value = '=HYPERLINK("http://evil.example","click")'
    rec.cells["usvp"].value = "bad\x0bchar\x0c here"
    rec.sources[0].title = "@SUM(A1:A9)"
    path = excel.export(p, tmp_path / "safe.xlsx")           # must not raise
    wb = load_workbook(path)
    cells = [c for ws in wb.worksheets for row in ws.iter_rows() for c in row]
    assert not any(c.data_type == "f" for c in cells)
    ca = wb["Company Analysis"]
    desc = next(r for r in range(4, 60) if ca.cell(r, 2).value == "Description")
    assert ca.cell(desc, 4).value.startswith("=HYPERLINK")    # kept as visible text, not executed
    usvp = next(r for r in range(4, 60) if ca.cell(r, 2).value.startswith("Unique Selling"))
    assert ca.cell(usvp, 4).value == "badchar here"


def test_scores_never_land_on_the_wrong_feature(corpus):
    """Reviewer's case: ids '1','2' with AI omitted -> AI must be '?', not Email marketing's score."""
    class Shifted(FakeLLM):
        def generate_json(self, prompt, schema, system=""):
            mk = lambda i, name, sc: {"id": i, "feature": name, "standard": True, "premium": None, "enterprise": None,
                                     "score": sc, "rationale": "r", "confidence": "high",
                                     "citations": [{"source_id": "S1", "quote": "build interactive 3D floor plans"}]}
            data = {"scores": [mk("1", "Email marketing", 4), mk("F03", "Reporting", 3), mk("F3", "Reporting", 5)]}
            return schema.model_validate_json(json.dumps(data))
    feats = [FeatureDef(category="X", name="AI"), FeatureDef(category="X", name="Email marketing"),
             FeatureDef(category="X", name="Reporting")]
    out = {f.feature: f.score for f in score.score(_record(corpus), feats, Intake(client_name="C", product="p"), Shifted())}
    assert out == {"AI": None, "Email marketing": 4.0, "Reporting": 3.0}


def test_rejected_content_never_reaches_deliverables(tmp_path, corpus):
    p = _full_project(corpus)
    p.candidates = [Candidate(name="Acme", status="accepted"), Candidate(name="Beta", status="rejected")]
    p.competitors[0].swot["strengths"].status = "rejected"
    p.competitors[0].swot["strengths"].value = "SECRET WRONG STRENGTH"
    wb = load_workbook(excel.export(p, tmp_path / "x.xlsx"))
    assert [c.value for c in wb["Company Analysis"][2]][3:] == ["Acme"]            # Beta rejected after research
    texts = lambda prs: " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    deck = texts(Presentation(pptx.export_deck(p, tmp_path / "d.pptx")))
    cards = texts(Presentation(pptx.export_battlecards(p, tmp_path / "c.pptx")))
    assert "Beta" not in deck and "Beta" not in cards
    assert "SECRET WRONG STRENGTH" not in deck and "SECRET WRONG STRENGTH" not in cards


def test_pilot_headcount_ignores_years_and_other_numbers():
    from ca_tool.pilot import fact_agrees as f
    assert f("employees", "1,000-5,000", "~200 employees (LinkedIn, 2024)") is False
    assert f("employees", "51-200", "1,200 employees across 3 offices") is False
    assert f("employees", "51-200", "About 120 employees (2025)") is True
    assert f("employees", "3,601", "3,500 employees") is True


def _fake_client(behaviour):
    """behaviour(model, n) -> text to return, or an exception to raise."""
    from google.genai import errors
    calls = []

    class Models:
        def generate_content(self, model, contents, config):
            calls.append(model)
            out = behaviour(model, len(calls))
            if isinstance(out, Exception):
                raise out
            part = type("P", (), {"text": out})()
            cand = type("C", (), {"content": type("X", (), {"parts": [part]})()})()
            return type("R", (), {"candidates": [cand]})()
    return type("Cl", (), {"models": Models()})(), calls, errors


def test_llm_error_handling():
    from pydantic import BaseModel
    from ca_tool.llm import gemini as g
    from ca_tool.llm.base import LLMError

    class Out(BaseModel):
        x: int

    g.GeminiLLM._cooldown.clear()
    # 400 on m1 -> next model is tried and m1 is NOT benched for later calls
    client, calls, errors = _fake_client(lambda m, n: errors.APIError(400, {"error": {"code": 400, "message": "bad", "status": "INVALID_ARGUMENT"}}) if m == "m1" else '{"x": 1}')
    llm = g.GeminiLLM(api_key="k", models=["m1", "m2"], retries=2, round_wait=0)
    llm.client = client
    assert llm.generate_json("p", Out).x == 1 and "m1" not in g.GeminiLLM._cooldown
    # malformed JSON twice -> LLMError (not a pydantic traceback)
    client, calls, _ = _fake_client(lambda m, n: "not json")
    llm.client = client
    with pytest.raises(LLMError):
        llm.generate_json("p", Out)
    # empty reply: each model asked once per pass, no waiting round
    client, calls, _ = _fake_client(lambda m, n: "")
    llm.client = client
    with pytest.raises(LLMError):
        llm.generate_text("p")
    assert calls == ["m1", "m2"]
    g.GeminiLLM._cooldown.clear()


def test_gate1_merge_handles_new_rows_renames_and_repointing(tmp_path, monkeypatch, corpus):
    import importlib, sys, types
    import pandas as pd
    # import merge_candidates from app.py without running the Streamlit script
    src = open("app.py", encoding="utf-8").read()
    start, end = src.index("STATUSES = ("), src.index("def bump(")
    ns = {"pd": pd, "Candidate": Candidate, "Project": Project}
    exec(src[start:end], ns)
    merge = ns["merge_candidates"]
    p = _full_project(corpus)
    p.candidates = [Candidate(name="Acme", website="https://acme.example", status="accepted",
                              citations=[Citation(source_id="D1", quote="q", verified=True)]),
                    Candidate(name="Beta", website="https://beta.example", status="accepted")]
    a, b = p.candidates
    df = pd.DataFrame([{"_id": a.uid, "status": "accepted", "name": "Acme Spaces", "website": "https://acme.example", "type": "direct"},
                       {"_id": b.uid, "status": "accepted", "name": "Beta", "website": "https://beta-real.example", "type": "direct"}])
    new_row = pd.DataFrame([{"_id": None, "status": None, "name": "Gamma", "website": float("nan"), "type": None}], index=[2])
    out = merge(p, pd.concat([df, new_row]))
    assert [c.name for c in out] == ["Acme Spaces", "Beta", "Gamma"]
    assert out[0].citations and out[0].status == "edited"            # rename kept evidence, marked edited
    assert out[2].status == "accepted" and out[2].website == ""      # added row: valid defaults, no None/NaN
    assert p.record("Acme") is None and p.record("Beta") is None      # renamed / re-pointed research discarded
    p.candidates = out
    Project.model_validate_json(p.model_dump_json())                 # project still loads



def test_gate1_delete_and_add_in_one_edit_does_not_transfer_evidence(corpus):
    """Reviewer regression: delete the last row and add a new one -> the new row re-used the
    deleted row's position and inherited its citations and approval."""
    import pandas as pd
    src = open("app.py", encoding="utf-8").read()
    ns = {"pd": pd, "Candidate": Candidate, "Project": Project}
    exec(src[src.index("STATUSES = ("): src.index("def apply_score_edits(")], ns)
    p = Project(id="t", intake=Intake(client_name="C", product="p"))
    p.candidates = [Candidate(name="A", status="accepted"),
                    Candidate(name="Foo", status="accepted", citations=[Citation(source_id="D1", quote="Foo is great", verified=True)])]
    # Streamlit: Foo deleted, Bar appended at the freed position 1, with no id
    df = pd.DataFrame([{"_id": p.candidates[0].uid, "status": "accepted", "name": "A"},
                       {"_id": None, "status": None, "name": "Bar"}])
    out = ns["merge_candidates"](p, df)
    bar = next(c for c in out if c.name == "Bar")
    assert bar.citations == [] and bar.status == "accepted" and bar.uid != p.candidates[1].uid



def test_quote_checker_second_review_cases():
    page = ("Acme is the choice of many venues. It is not the market leader in Europe. Its plan does include SSO. "
            "The company raised $5 million in 2019. Headquartered in Boston. " + "Filler text about other things. " * 20 +
            "Acme also sells hotel event software.")
    assert not quote_in_text("Its plan does not include SSO", page)                                   # negation flip
    assert not quote_in_text("The company raised $50 million in 2019", page)                         # number flip
    assert not quote_in_text("Acme is the ... market leader in Europe ... hotel event software", page)  # far-apart stitching
    assert quote_in_text("It is not the market leader in Europe", page)
    assert quote_in_text("Acme is the choice ... not the market leader in Europe", page)             # close parts ok
    assert quote_in_text("Headquartered in Boston", page)                                            # 3-word fact
    from ca_tool.research.corpus import snippet_url
    line = "- Acme - Wikipedia: Acme is a venue software company based in Boston (https://en.wikipedia.org/wiki/Acme_(company))"
    assert snippet_url(line, "Acme is a venue software company") == "https://en.wikipedia.org/wiki/Acme_(company)"


def test_score_ids_trusted_when_the_model_rewords_names(corpus):
    class Reworded(FakeLLM):
        def generate_json(self, prompt, schema, system=""):
            mk = lambda i, name, sc: {"id": i, "feature": name, "standard": True, "premium": None, "enterprise": None,
                                     "score": sc, "rationale": "r", "confidence": "high",
                                     "citations": [{"source_id": "S1", "quote": "build interactive 3D floor plans"}]}
            data = {"scores": [mk("F1", "Artificial intelligence", 2), mk("F2", "Email campaigns", 4), mk("F3", "Reports", 3)]}
            return schema.model_validate_json(json.dumps(data))
    feats = [FeatureDef(category="X", name="AI"), FeatureDef(category="X", name="Email marketing"),
             FeatureDef(category="X", name="Reporting")]
    out = {f.feature: f.score for f in score.score(_record(corpus), feats, Intake(client_name="C", product="p"), Reworded())}
    assert out == {"AI": 2.0, "Email marketing": 4.0, "Reporting": 3.0}


def test_pilot_headcount_ranges_with_year_like_numbers():
    from ca_tool.pilot import _emp_range
    assert _emp_range("1001-2000 employees") == (1001, 2000)
    assert _emp_range("2,000-5,000") == (2000, 5000)
    assert _emp_range("~200 employees (LinkedIn, 2024)") == (200, 200)
